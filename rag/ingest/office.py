from __future__ import annotations
from pathlib import Path
from docx import Document as DocxDocument
from pptx import Presentation

def extract_docx_text(path: Path) -> str:
    d = DocxDocument(str(path))
    parts = [p.text.strip() for p in d.paragraphs if p.text and p.text.strip()]
    return "\n".join(parts).strip()

def extract_pptx_text(path: Path) -> str:
    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_lines = [f"# Slide {i}"]
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text and shape.text.strip():
                slide_lines.append(shape.text.strip())
        out.append("\n".join(slide_lines))
    return "\n\n".join(out).strip()
